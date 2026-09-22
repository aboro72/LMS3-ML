from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "LASTENHEFT.md"
OUTPUT = ROOT / "lastenheft.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F4F7FA"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRID, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tbl_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tbl_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tbl_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=BLACK, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(clean_inline(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    borders(cell)


def clean_inline(text):
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("**", "")
    text = text.replace("`", "")
    return text


def add_rich_paragraph(doc, text, style=None, bullet=False):
    p = doc.add_paragraph(style=style)
    if bullet:
        p.style = "List Bullet"
    p.paragraph_format.widow_control = True
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run(clean_inline(part.replace("**", "").replace("`", "")))
        if part.startswith("**") and part.endswith("**"):
            run.bold = True
        if part.startswith("`") and part.endswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    return p


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, space_before, space_after in (
        ("Heading 1", 15, 18, 8),
        ("Heading 2", 12, 13, 6),
        ("Heading 3", 10.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for style_name in ("Title", "Subtitle"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = BLACK
        p_pr = style._element.get_or_add_pPr()
        for child in list(p_pr):
            if child.tag == qn("w:pBdr"):
                p_pr.remove(child)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ABoroLMS  |  Lastenheft")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Vertrauliche Projektunterlage  •  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Lastenheft")
    run.font.size = Pt(31)
    run.font.bold = True
    run.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("ABoroLMS und ML Gruppe")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(64, 64, 64)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Fachliche Anforderungen und belastbarer aktueller Projektstand")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(4.0), Cm(9.6)]
    values = [
        ("Dokumentstand", "21.09.2026"),
        ("Bewertete Revision", "f4fa398"),
        ("Zielplattform", "Selfhosted Django LMS"),
        ("Betreiberzielbild", "ML Gruppe, Einzelinstallation"),
        ("Reifegrad", "Entwicklungsstand, nicht produktionsabnahmefähig"),
    ]
    for row, (label, value) in zip(table.rows, values):
        row.cells[0].width = widths[0]
        row.cells[1].width = widths[1]
        set_cell_text(row.cells[0], label, bold=True, size=9)
        set_cell_text(row.cells[1], value, size=9)
        shade(row.cells[0], BLUE)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run("Dieses Dokument trennt nachprüfbaren Ist-Stand, Zielanforderungen und offene Abnahmepunkte.")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_page_break()


def add_contents_page(doc, headings):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Inhaltsübersicht")
    doc.add_paragraph("Die Abschnittsnummern entsprechen dem Lastenheft. Tabellen und Anforderungscodes dienen der späteren Abnahme.")
    for level, text in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5 * (level - 1))
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(clean_inline(text))
        run.font.size = Pt(10 if level == 2 else 9)
        run.bold = level == 2
    doc.add_page_break()


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    repeat_header(header)
    for i, value in enumerate(rows[0]):
        set_cell_text(header.cells[i], value, bold=True, color=BLACK, size=8.2)
        shade(header.cells[i], BLUE)
    for row_index, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=8.0)
            if row_index % 2 == 0:
                shade(cells[i], PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = []
    for line in lines:
        match = re.match(r"^(#{2,3})\s+(.+)$", line)
        if match:
            headings.append((len(match.group(1)), match.group(2)))

    doc = Document()
    set_document_defaults(doc)
    add_title_page(doc)
    add_contents_page(doc, headings)

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.4)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                for idx, code in enumerate(code_lines):
                    run = p.add_run(code + ("\n" if idx < len(code_lines) - 1 else ""))
                    run.font.name = "Consolas"
                    run.font.size = Pt(8.5)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("# "):
            i += 1
            continue
        match = re.match(r"^(#{2,3})\s+(.+)$", line)
        if match:
            level = len(match.group(1)) - 1
            doc.add_heading(clean_inline(match.group(2)), level=level)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|") and re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", lines[i + 1]):
            table_rows = []
            table_rows.append([x.strip() for x in line.strip().strip("|").split("|")])
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, table_rows)
            continue
        if line.startswith("- "):
            add_rich_paragraph(doc, line[2:], bullet=True)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            add_rich_paragraph(doc, re.sub(r"^\d+\.\s+", "", line), style="List Number")
            i += 1
            continue
        add_rich_paragraph(doc, line)
        i += 1

    doc.core_properties.title = "Lastenheft ABoroLMS und ML Gruppe"
    doc.core_properties.subject = "Fachliche Anforderungen und aktueller Projektstand"
    doc.core_properties.author = "ABoroLMS Projekt"
    doc.core_properties.keywords = "LMS, Lastenheft, ML Gruppe, Prüfungen, Zertifikate"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
